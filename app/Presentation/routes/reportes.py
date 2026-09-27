import json
from io import BytesIO

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.Presentation.schemas.reporte import (
    ReporteResponse, ReporteGeneradoResponse, ReporteContenidoResponse
)
from app.Presentation.dependencies import require_roles, get_service, GESTION
from typing import List

router = APIRouter(prefix="/reportes", tags=["reportes"])

HEADER_FILL = PatternFill("solid", fgColor="3498DB")
HEADER_FONT = Font(color="FFFFFF", bold=True)


@router.get("/", response_model=List[ReporteResponse])
def listar_reportes(
    use_cases=Depends(get_service),
    current_user=Depends(require_roles(*GESTION)),
):
    return use_cases["listar_reportes"].execute()


@router.get("/{id}", response_model=ReporteContenidoResponse)
def obtener_reporte(
    id: int,
    use_cases=Depends(get_service),
    current_user=Depends(require_roles(*GESTION)),
):
    reporte = use_cases["obtener_reporte"].execute(id)
    return {
        "id_reporte": reporte.id_reporte,
        "tipo_reporte": reporte.tipo_reporte,
        "periodo": reporte.periodo,
        "fecha_generacion": reporte.fecha_generacion,
        "datos": use_cases["obtener_reporte"].parsear_contenido(reporte),
    }


@router.post("/diario", response_model=ReporteGeneradoResponse, status_code=201)
def generar_reporte_diario(
    use_cases=Depends(get_service),
    current_user=Depends(require_roles(*GESTION)),
):
    resultado = use_cases["generar_reporte_diario"].execute()
    return {"message": "Reporte diario generado", "id_reporte": resultado["id_reporte"]}


@router.post("/mensual", response_model=ReporteGeneradoResponse, status_code=201)
def generar_reporte_mensual(
    use_cases=Depends(get_service),
    current_user=Depends(require_roles(*GESTION)),
):
    resultado = use_cases["generar_reporte_mensual"].execute()
    return {"message": "Reporte mensual generado", "id_reporte": resultado["id_reporte"]}


@router.get("/{id}/excel")
def descargar_reporte_excel(
    id: int,
    use_cases=Depends(get_service),
    current_user=Depends(require_roles(*GESTION)),
):
    use_case = use_cases["obtener_reporte"]
    reporte = use_case.execute(id)
    datos = use_case.parsear_contenido(reporte)

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"

    ws.cell(row=1, column=1, value="Reporte").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value="Tipo").font = Font(bold=True)
    ws.cell(row=2, column=2, value=reporte.tipo_reporte)
    ws.cell(row=3, column=1, value="Periodo").font = Font(bold=True)
    ws.cell(row=3, column=2, value=reporte.periodo)
    ws.cell(row=4, column=1, value="Generado").font = Font(bold=True)
    ws.cell(row=4, column=2, value=str(reporte.fecha_generacion))

    header_row = 6
    ws.cell(row=header_row, column=1, value="Métrica").fill = HEADER_FILL
    ws.cell(row=header_row, column=1).font = HEADER_FONT
    ws.cell(row=header_row, column=2, value="Valor").fill = HEADER_FILL
    ws.cell(row=header_row, column=2).font = HEADER_FONT
    ws.cell(row=header_row, column=1).alignment = Alignment(horizontal="center")

    for offset, (clave, valor) in enumerate(datos.items(), start=1):
        fila = header_row + offset
        ws.cell(row=fila, column=1, value=clave)
        ws.cell(row=fila, column=2, value=valor)

    ws.column_dimensions[get_column_letter(1)].width = 32
    ws.column_dimensions[get_column_letter(2)].width = 24

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"reporte_{reporte.tipo_reporte}_{reporte.periodo}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
